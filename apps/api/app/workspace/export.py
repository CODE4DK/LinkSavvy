"""Single-asset export as TXT/MD/PDF/DOCX, and a whole selection as a
ZIP -- plain single-column documents, same simple approach as
app.career.export (no tables, no styling worth breaking across
formats).
"""

from __future__ import annotations

import io
import zipfile

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from app.models.asset import Asset

_PAGE_WIDTH, _PAGE_HEIGHT = LETTER
_MARGIN = 60
_BODY_SIZE = 11
_TITLE_SIZE = 16
_LINE_HEIGHT = 15
_FONT = "Helvetica"
_FONT_BOLD = "Helvetica-Bold"
_WRAP_WIDTH = 95


def _wrapped_lines(text: str) -> list[str]:
    lines: list[str] = []
    for paragraph in text.splitlines() or [""]:
        if not paragraph:
            lines.append("")
            continue
        words = paragraph.split(" ")
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > _WRAP_WIDTH and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def export_asset_txt(asset: Asset) -> bytes:
    return f"{asset.title}\n\n{asset.body}\n".encode()


def export_asset_md(asset: Asset) -> bytes:
    return f"# {asset.title}\n\n{asset.body}\n".encode()


def export_asset_pdf(asset: Asset) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    y = _PAGE_HEIGHT - _MARGIN

    def _new_page_if_needed() -> None:
        nonlocal y
        if y < _MARGIN:
            pdf.showPage()
            pdf.setFont(_FONT, _BODY_SIZE)
            y = _PAGE_HEIGHT - _MARGIN

    pdf.setFont(_FONT_BOLD, _TITLE_SIZE)
    pdf.drawString(_MARGIN, y, asset.title)
    y -= _LINE_HEIGHT * 1.5
    pdf.setFont(_FONT, _BODY_SIZE)

    for line in _wrapped_lines(asset.body):
        _new_page_if_needed()
        pdf.drawString(_MARGIN, y, line)
        y -= _LINE_HEIGHT

    pdf.save()
    return buffer.getvalue()


def export_asset_docx(asset: Asset) -> bytes:
    docx = DocxDocument()
    docx.add_heading(asset.title, level=1)
    for paragraph in asset.body.splitlines() or [""]:
        docx.add_paragraph(paragraph)
    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


_EXPORTERS = {
    "txt": export_asset_txt,
    "md": export_asset_md,
    "pdf": export_asset_pdf,
    "docx": export_asset_docx,
}
_EXTENSIONS = {"txt": ".txt", "md": ".md", "pdf": ".pdf", "docx": ".docx"}


def export_asset(asset: Asset, *, fmt: str) -> bytes:
    exporter = _EXPORTERS.get(fmt)
    if exporter is None:
        raise ValueError(f"unsupported export format: {fmt!r}")
    return exporter(asset)


def export_assets_zip(assets: list[Asset], *, fmt: str = "txt") -> bytes:
    buffer = io.BytesIO()
    extension = _EXTENSIONS[fmt]
    seen_names: dict[str, int] = {}
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for asset in assets:
            base_name = "".join(
                c if c.isalnum() or c in " -_" else "_" for c in asset.title
            ).strip()
            base_name = base_name or "asset"
            count = seen_names.get(base_name, 0)
            seen_names[base_name] = count + 1
            filename = (
                f"{base_name}{extension}" if count == 0 else f"{base_name} ({count}){extension}"
            )
            zf.writestr(filename, export_asset(asset, fmt=fmt))
    return buffer.getvalue()
