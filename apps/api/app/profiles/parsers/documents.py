"""Text extraction for uploaded resumes (PDF, DOCX).

Extracted text is handed to the same deterministic parser used for the
paste path (`app.profiles.parsers.paste`) — there is no separate
"upload parser", only a separate way of getting text.
"""

from __future__ import annotations

import io

import pypdfium2 as pdfium
from docx import Document as DocxDocument
from pdfplumber.pdf import PDF as PdfplumberPDF

PDF_MAGIC = b"%PDF-"
DOCX_MAGIC = b"PK\x03\x04"


class UnsupportedDocumentError(Exception):
    """Raised with a message safe to show directly to the user."""


def sniff_upload_kind(data: bytes, content_type: str) -> str:
    """Returns 'pdf' or 'docx', validating magic bytes against the
    declared content type rather than trusting either alone."""
    is_pdf_type = content_type == "application/pdf"
    is_docx_type = content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    if data.startswith(PDF_MAGIC):
        if not is_pdf_type:
            raise UnsupportedDocumentError(
                "This file's content doesn't match its declared type. Please upload a "
                "genuine PDF or DOCX file."
            )
        return "pdf"

    if data.startswith(DOCX_MAGIC):
        if not is_docx_type:
            raise UnsupportedDocumentError(
                "This file's content doesn't match its declared type. Please upload a "
                "genuine PDF or DOCX file."
            )
        return "docx"

    raise UnsupportedDocumentError(
        "We couldn't recognise this file as a PDF or DOCX. Please upload one of those, "
        "or use the paste option instead."
    )


def extract_pdf_text(data: bytes) -> str:
    text = _extract_pdf_text_pdfplumber(data)
    if text is None:
        text = _extract_pdf_text_pdfium(data)
    if not text.strip():
        raise UnsupportedDocumentError(
            "We couldn't find any text in this PDF — it may be a scanned image. "
            "Please try the paste option or fill in your profile manually."
        )
    return text


def _extract_pdf_text_pdfplumber(data: bytes) -> str | None:
    try:
        pdf = PdfplumberPDF.open(io.BytesIO(data))
    except Exception:
        return None
    try:
        if getattr(pdf.doc, "is_extractable", True) is False:
            raise UnsupportedDocumentError(
                "This PDF is encrypted and can't be read. Please remove the password, "
                "or use the paste option instead."
            )
        pages_text = [page.extract_text() or "" for page in pdf.pages]
        return "\n\n".join(pages_text)
    except UnsupportedDocumentError:
        raise
    except Exception:
        return None
    finally:
        pdf.close()


def _extract_pdf_text_pdfium(data: bytes) -> str:
    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        message = str(exc).lower()
        if "password" in message or "encrypt" in message:
            raise UnsupportedDocumentError(
                "This PDF is encrypted and can't be read. Please remove the password, "
                "or use the paste option instead."
            ) from exc
        raise UnsupportedDocumentError(
            "This PDF couldn't be read — it may be corrupted. Please try the paste "
            "option or fill in your profile manually."
        ) from exc

    pages_text = []
    for page in pdf:
        text_page = page.get_textpage()
        pages_text.append(text_page.get_text_range())
    return "\n\n".join(pages_text)


def extract_docx_text(data: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocumentError(
            "This DOCX file couldn't be read — it may be corrupted. Please try the "
            "paste option or fill in your profile manually."
        ) from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        is_bullet = "bullet" in (paragraph.style.name or "").lower() if paragraph.style else False
        lines.append(f"• {text}" if is_bullet else text)

    extracted = "\n".join(lines)
    if not extracted.strip():
        raise UnsupportedDocumentError(
            "We couldn't find any text in this document. Please try the paste option "
            "or fill in your profile manually."
        )
    return extracted


def extract_text(data: bytes, content_type: str) -> tuple[str, str]:
    """Returns (kind, extracted_text) where kind is 'pdf' or 'docx'."""
    kind = sniff_upload_kind(data, content_type)
    if kind == "pdf":
        return kind, extract_pdf_text(data)
    return kind, extract_docx_text(data)
