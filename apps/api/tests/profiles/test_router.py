import io

from docx import Document as DocxDocument
from httpx import AsyncClient


async def _access_token(client: AsyncClient, creds: dict[str, str]) -> str:
    login = await client.post("/api/v1/auth/login", json=creds)
    return str(login.json()["access_token"])


async def _auth_headers(client: AsyncClient, creds: dict[str, str]) -> dict[str, str]:
    token = await _access_token(client, creds)
    return {"Authorization": f"Bearer {token}"}


MINIMAL_SNAPSHOT = {
    "source": "manual",
    "captured_at": "2026-01-01T00:00:00Z",
    "identity": {"full_name": "Jane Doe", "headline": "Engineer"},
}


async def test_profile_endpoints_require_auth(client: AsyncClient) -> None:
    response = await client.get("/api/v1/profile/snapshot")
    assert response.status_code == 401


async def test_put_snapshot_commits_manual_source(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.put("/api/v1/profile/snapshot", json=MINIMAL_SNAPSHOT, headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["version"] == 1
    assert body["source"] == "manual"
    assert body["is_active"] is True
    assert isinstance(body["completeness_score"], int)


async def test_put_snapshot_ignores_client_supplied_source(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    spoofed = {**MINIMAL_SNAPSHOT, "source": "linkedin_api"}
    response = await client.put("/api/v1/profile/snapshot", json=spoofed, headers=headers)
    assert response.status_code == 200
    assert response.json()["source"] == "manual"


async def test_get_active_snapshot_returns_full_payload(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    await client.put("/api/v1/profile/snapshot", json=MINIMAL_SNAPSHOT, headers=headers)

    response = await client.get("/api/v1/profile/snapshot", headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["payload"]["identity"]["full_name"] == "Jane Doe"


async def test_get_active_snapshot_404_before_any_commit(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.get("/api/v1/profile/snapshot", headers=headers)
    assert response.status_code == 404


async def test_second_commit_bumps_version_and_deactivates_first(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    first = await client.put("/api/v1/profile/snapshot", json=MINIMAL_SNAPSHOT, headers=headers)
    second_payload = {
        **MINIMAL_SNAPSHOT,
        "identity": {"full_name": "Jane Doe", "headline": "Senior Engineer"},
    }
    second = await client.put("/api/v1/profile/snapshot", json=second_payload, headers=headers)
    assert first.json()["version"] == 1
    assert second.json()["version"] == 2

    listing = await client.get("/api/v1/profile/snapshots", headers=headers)
    versions = {row["version"]: row["is_active"] for row in listing.json()}
    assert versions == {1: False, 2: True}


async def test_paste_import_then_commit_flow(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    paste_text = "Jane Doe\nSenior Engineer\n\nAbout\nBuilds things. Ships things. Fixes things.\n"
    import_response = await client.post(
        "/api/v1/profile/imports", json={"text": paste_text}, headers=headers
    )
    assert import_response.status_code == 200
    body = import_response.json()
    assert body["status"] == "needs_review"
    assert body["draft"]["identity"]["full_name"] == "Jane Doe"
    import_id = body["import_id"]

    commit_response = await client.post(
        f"/api/v1/profile/imports/{import_id}/commit",
        json={"payload": body["draft"]},
        headers=headers,
    )
    assert commit_response.status_code == 200
    assert commit_response.json()["source"] == "paste"
    assert commit_response.json()["version"] == 1


async def test_paste_import_rejects_empty_text(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post("/api/v1/profile/imports", json={"text": "   "}, headers=headers)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_FAILED"


async def test_committing_an_import_twice_fails(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    import_response = await client.post(
        "/api/v1/profile/imports", json={"text": "Jane Doe\nEngineer\n"}, headers=headers
    )
    import_id = import_response.json()["import_id"]
    draft = import_response.json()["draft"]

    first_commit = await client.post(
        f"/api/v1/profile/imports/{import_id}/commit", json={"payload": draft}, headers=headers
    )
    assert first_commit.status_code == 200

    second_commit = await client.post(
        f"/api/v1/profile/imports/{import_id}/commit", json={"payload": draft}, headers=headers
    )
    assert second_commit.status_code == 422


async def test_upload_docx_import_flow(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    document = DocxDocument()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Senior Engineer")
    buffer = io.BytesIO()
    document.save(buffer)

    response = await client.post(
        "/api/v1/profile/imports/upload",
        headers=headers,
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "needs_review"
    assert body["draft"]["identity"]["full_name"] == "Jane Doe"


async def test_upload_rejects_unrecognised_file(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post(
        "/api/v1/profile/imports/upload",
        headers=headers,
        files={"file": ("notes.txt", b"just some text", "text/plain")},
    )
    assert response.status_code == 422


async def test_diff_between_two_versions(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    await client.put("/api/v1/profile/snapshot", json=MINIMAL_SNAPSHOT, headers=headers)
    second_payload = {
        "source": "manual",
        "captured_at": "2026-01-02T00:00:00Z",
        "identity": {"full_name": "Jane Doe", "headline": "Staff Engineer"},
        "skills": [{"name": "Python"}],
    }
    await client.put("/api/v1/profile/snapshot", json=second_payload, headers=headers)

    response = await client.get("/api/v1/profile/snapshots/1/diff/2", headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["from_version"] == 1
    assert body["to_version"] == 2

    headline_change = next(fc for fc in body["field_changes"] if fc["path"] == "/identity/headline")
    assert headline_change["before"] == "Engineer"
    assert headline_change["after"] == "Staff Engineer"

    assert body["list_changes"]["skills"][0]["change"] == "added"
    assert body["list_changes"]["skills"][0]["key"] == "Python"


async def test_diff_missing_version_404s(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    await client.put("/api/v1/profile/snapshot", json=MINIMAL_SNAPSHOT, headers=headers)
    response = await client.get("/api/v1/profile/snapshots/1/diff/99", headers=headers)
    assert response.status_code == 404


async def test_sync_without_connection_returns_not_found(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post("/api/v1/profile/sync", headers=headers)
    assert response.status_code == 404


async def test_connect_requires_configured_linkedin(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post("/api/v1/profile/connect", headers=headers)
    assert response.status_code == 500
