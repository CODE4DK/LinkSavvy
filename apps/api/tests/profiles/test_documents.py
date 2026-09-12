import io

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.profiles.parsers.documents import (
    UnsupportedDocumentError,
    extract_docx_text,
    extract_pdf_text,
    extract_text,
    sniff_upload_kind,
)

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _build_pdf_bytes(lines: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _build_docx_bytes(paragraphs: list[str], bullets: list[str]) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    for text in bullets:
        document.add_paragraph(text, style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_sniff_upload_kind_detects_pdf_by_magic_bytes() -> None:
    data = _build_pdf_bytes(["Hello"])
    assert sniff_upload_kind(data, PDF_CONTENT_TYPE) == "pdf"


def test_sniff_upload_kind_detects_docx_by_magic_bytes() -> None:
    data = _build_docx_bytes(["Hello"], [])
    assert sniff_upload_kind(data, DOCX_CONTENT_TYPE) == "docx"


def test_sniff_upload_kind_rejects_mismatched_content_type() -> None:
    data = _build_pdf_bytes(["Hello"])
    with pytest.raises(UnsupportedDocumentError):
        sniff_upload_kind(data, DOCX_CONTENT_TYPE)


def test_sniff_upload_kind_rejects_unrecognised_file() -> None:
    with pytest.raises(UnsupportedDocumentError):
        sniff_upload_kind(b"not a real document", PDF_CONTENT_TYPE)


def test_extract_pdf_text_finds_real_text() -> None:
    data = _build_pdf_bytes(["Jane Doe", "Senior Engineer"])
    text = extract_pdf_text(data)
    assert "Jane Doe" in text
    assert "Senior Engineer" in text


def test_extract_pdf_text_raises_on_empty_pdf() -> None:
    data = _build_pdf_bytes([])
    with pytest.raises(UnsupportedDocumentError, match="scanned image"):
        extract_pdf_text(data)


def test_extract_docx_text_preserves_paragraphs_and_bullets() -> None:
    data = _build_docx_bytes(["Jane Doe", "About me"], ["Did a thing", "Did another thing"])
    text = extract_docx_text(data)
    lines = text.split("\n")
    assert "Jane Doe" in lines
    assert "About me" in lines
    assert "• Did a thing" in lines
    assert "• Did another thing" in lines


def test_extract_docx_text_raises_on_empty_document() -> None:
    data = _build_docx_bytes([], [])
    with pytest.raises(UnsupportedDocumentError):
        extract_docx_text(data)


def test_extract_docx_text_raises_on_corrupt_file() -> None:
    with pytest.raises(UnsupportedDocumentError):
        extract_docx_text(b"PK\x03\x04not a real docx")


def test_extract_text_dispatches_by_kind() -> None:
    pdf_data = _build_pdf_bytes(["Hello PDF"])
    kind, text = extract_text(pdf_data, PDF_CONTENT_TYPE)
    assert kind == "pdf"
    assert "Hello PDF" in text

    docx_data = _build_docx_bytes(["Hello DOCX"], [])
    kind, text = extract_text(docx_data, DOCX_CONTENT_TYPE)
    assert kind == "docx"
    assert "Hello DOCX" in text
