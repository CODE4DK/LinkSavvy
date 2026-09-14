from __future__ import annotations

import io

from docx import Document as DocxDocument
from httpx import AsyncClient

from app.content.voice_service import MIN_SAMPLES


async def _access_token(client: AsyncClient, creds: dict[str, str]) -> str:
    login = await client.post("/api/v1/auth/login", json=creds)
    return str(login.json()["access_token"])


async def _auth_headers(client: AsyncClient, creds: dict[str, str]) -> dict[str, str]:
    token = await _access_token(client, creds)
    return {"Authorization": f"Bearer {token}"}


def _sample_texts(count: int) -> list[str]:
    return [f"This is sample post number {i} about my work in engineering." for i in range(count)]


async def test_voice_endpoints_require_auth(client: AsyncClient) -> None:
    response = await client.get("/api/v1/content/voice")
    assert response.status_code == 401


async def test_get_voice_profile_defaults_before_any_submission(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.get("/api/v1/content/voice", headers=headers)
    assert response.status_code == 200
    body = response.json()
    assert body["source"] == "default"
    assert body["sample_count"] == 0
    assert "clear" in body["tone_adjectives"]


async def test_submit_pasted_samples_derives_a_descriptor(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post(
        "/api/v1/content/voice/samples",
        headers=headers,
        json={"texts": _sample_texts(MIN_SAMPLES)},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["source"] == "derived"
    assert body["sample_count"] == MIN_SAMPLES
    assert body["tone_adjectives"]

    follow_up = await client.get("/api/v1/content/voice", headers=headers)
    assert follow_up.json()["source"] == "derived"


async def test_submit_pasted_samples_rejects_too_few(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post(
        "/api/v1/content/voice/samples",
        headers=headers,
        json={"texts": _sample_texts(MIN_SAMPLES - 1)},
    )
    assert response.status_code == 422


async def test_submit_uploaded_docx_splits_on_separator_and_derives(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    document = DocxDocument()
    for i, text in enumerate(_sample_texts(MIN_SAMPLES)):
        if i > 0:
            document.add_paragraph("---")
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)

    response = await client.post(
        "/api/v1/content/voice/samples/upload",
        headers=headers,
        files={
            "file": (
                "posts.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["source"] == "derived"
    assert body["sample_count"] == MIN_SAMPLES


async def test_upload_rejects_unrecognised_file(
    client: AsyncClient, registered_user: dict[str, str]
) -> None:
    headers = await _auth_headers(client, registered_user)
    response = await client.post(
        "/api/v1/content/voice/samples/upload",
        headers=headers,
        files={"file": ("notes.txt", b"just some text", "text/plain")},
    )
    assert response.status_code == 422
